"""CSV/XLSX identity directives, native conversion and meaningful loss checks."""
import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from cmccdb_schema import dataset_constructor, message_helpers, spreadsheet
from cmccdb_schema.proto import dataset_pb2
from cmccdb_schema.scripts import construct_dataset
from cmccdb_schema.spreadsheet_test import dataset

DATASET_ID = "cmcc_dataset-" + "a" * 32
REACTION_ID = "cmcc-" + "b" * 32


def rows():
    # Deliberately ragged, with no common REACTION values. Blank IDs are generated.
    return [["DATASET ID", DATASET_ID], ["REACTION"], ["VARIANTS"],
            ["", "REACTION ID", "notes"], ["", "", "procedure_details"],
            ["DATA"], ["", REACTION_ID, "given ID"], ["", "", "generated ID"]]


class IdentityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def csv_file(self, records=None):
        path = self.root / "dataset.csv"
        with path.open("w", newline="", encoding="utf-8") as stream:
            csv.writer(stream).writerows(rows() if records is None else records)
        return path

    def test_csv_explicit_and_generated_identity(self):
        actual = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(str(self.csv_file()))
        self.assertEqual(actual.dataset_id, DATASET_ID)
        self.assertEqual(actual.reactions[0].reaction_id, REACTION_ID)
        self.assertEqual(actual.reactions[1].reaction_id, "cmcc-" + "a" * 28 + "0002")
        self.assertEqual(actual.reactions[0].notes.procedure_details, "given ID")

    def test_csv_text_and_binary_buffers_keep_metadata_and_values(self):
        stream = io.StringIO()
        csv.writer(stream).writerows(rows())
        for source in (io.StringIO(stream.getvalue()), io.BytesIO(stream.getvalue().encode("utf-8"))):
            with self.subTest(source=type(source).__name__):
                metadata = {}
                blocks = dataset_constructor.DatasetConstructor.from_spreadsheet(source, suffix=".csv",
                                                                                 dataset_metadata=metadata)
                self.assertEqual(metadata["dataset_id"], DATASET_ID)
                self.assertEqual(blocks[0][1][0][0], REACTION_ID)

    def test_xlsx_identity_directive_and_variants_across_sheets(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "Identity"
        workbook.active.append(["DATASET ID", DATASET_ID])
        for index in range(2):
            sheet = workbook.create_sheet("Reactions" + str(index))
            for row in rows()[1:7]:
                if row[0] == "":
                    row = list(row)
                    if row[1] == REACTION_ID:
                        row[1] = "cmcc-" + f"{index + 1:032x}"
                sheet.append(row)
        path = self.root / "dataset.xlsx"
        workbook.save(path); workbook.close()
        actual = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(str(path))
        self.assertEqual(actual.dataset_id, DATASET_ID)
        self.assertEqual([r.reaction_id for r in actual.reactions], ["cmcc-" + f"{i:032x}" for i in (1, 2)])

    def test_identity_precedence_and_legacy_fallback(self):
        path = str(self.csv_file())
        actual = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(path, id="c" * 32,
                                                                             default_id="d" * 32)
        self.assertEqual(actual.dataset_id, "cmcc_dataset-" + "c" * 32)
        actual = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(path, default_id="d" * 32)
        self.assertEqual(actual.dataset_id, DATASET_ID)
        path = str(self.csv_file(rows()[1:]))
        actual = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(path, default_id="d" * 32)
        self.assertEqual(actual.dataset_id, "cmcc_dataset-" + "d" * 32)
        generated = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(path)
        self.assertRegex(generated.dataset_id, r"^cmcc_dataset-[0-9a-f]{32}$")

    def test_identity_directives_can_appear_between_blocks_or_in_data(self):
        records = rows()[1:]
        records.insert(6, ["DATASET ID", DATASET_ID])
        records.append(["DATASET ID", DATASET_ID])
        actual = dataset_constructor.DatasetConstructor.enumerate_spreadsheet(str(self.csv_file(records)))
        self.assertEqual(len(actual.reactions), 2)
        self.assertEqual(actual.dataset_id, DATASET_ID)

    def test_bad_or_conflicting_identity_is_rejected(self):
        for directive in (["DATASET ID"], ["DATASET ID", ""], ["DATASET ID", "wrong"],
                          ["DATASET ID", DATASET_ID, "extra"],
                          ["DATASET ID", "cmcc_dataset-" + "c" * 32]):
            with self.subTest(directive=directive):
                path = self.csv_file(rows() + [directive])
                with self.assertRaisesRegex(ValueError, "DATASET ID"):
                    dataset_constructor.DatasetConstructor.enumerate_spreadsheet(str(path))

    def test_conflicting_ids_on_separate_xlsx_sheets_are_rejected(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["DATASET ID", DATASET_ID])
        workbook.create_sheet().append(["DATASET ID", "cmcc_dataset-" + "c" * 32])
        path = self.root / "conflicting.xlsx"
        workbook.save(path); workbook.close()
        with self.assertRaisesRegex(ValueError, "Conflicting DATASET ID"):
            dataset_constructor.DatasetConstructor.enumerate_spreadsheet(str(path))

    def test_duplicate_and_auto_generated_reaction_id_collisions_are_rejected(self):
        for value in (REACTION_ID, "cmcc-" + "a" * 28 + "0001"):
            records = rows()
            records[-1][1] = value
            if value != REACTION_ID:
                records[-2][1] = ""  # First generated ID collides with second supplied ID.
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, "Duplicate REACTION ID"):
                    dataset_constructor.DatasetConstructor.enumerate_spreadsheet(str(self.csv_file(records)))

    def test_native_upload_cli_honors_dataset_id_instead_of_file_hash(self):
        path = self.csv_file()
        output = self.root / "output.pbtxt"
        actual = construct_dataset.main({"--data": str(path), "--output": str(output),
            "--name": "Fixture", "--email": "fixture@example.invalid", "--no-validate": True})
        self.assertEqual(actual.dataset_id, DATASET_ID)
        self.assertEqual(actual.reactions[0].reaction_id, REACTION_ID)
        self.assertEqual(message_helpers.load_message(str(output), dataset_pb2.Dataset), actual)


class CsvTests(unittest.TestCase):
    def test_native_csv_never_calls_xlsx_backend_and_preserves_populated_fields(self):
        source = dataset(6)
        before = source.SerializeToString(deterministic=True)
        with patch.object(spreadsheet, "write_workbook", side_effect=AssertionError("No XLSX")), \
             patch.object(spreadsheet.subprocess, "run", side_effect=AssertionError("No subprocess")):
            payload = spreadsheet.serialize_csv(source)
        records = list(csv.reader(io.StringIO(payload.decode("utf-8"))))
        self.assertTrue(any(row[:2] == ["DATASET ID", DATASET_ID] for row in records))
        self.assertTrue(any("REACTION ID" in row for row in records))
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "test.csv"; path.write_bytes(payload)
            actual = spreadsheet.verify_round_trip(source, path)
        self.assertEqual(actual.dataset_id, source.dataset_id)
        self.assertEqual({r.reaction_id: r for r in actual.reactions}, {r.reaction_id: r for r in source.reactions})
        self.assertEqual(before, source.SerializeToString(deterministic=True))

    def test_idless_search_result_does_not_invent_an_exported_dataset_id(self):
        source = dataset(); source.ClearField("dataset_id")
        payload = spreadsheet.serialize_csv(source)
        self.assertFalse(any(r[0] == "DATASET ID" for r in csv.reader(io.StringIO(payload.decode()))))
        self.assertEqual(source.dataset_id, "")

    def test_valueless_wrappers_and_entirely_empty_lists_can_be_omitted(self):
        source = dataset()
        source.reactions[0].conditions.temperature.measurements.add()
        source.reactions[0].provenance.record_modified.add()
        before = source.SerializeToString(deterministic=True)
        self.assertTrue(spreadsheet.serialize_csv(source))
        self.assertEqual(source.SerializeToString(deterministic=True), before)

    def test_empty_oneof_choices_and_mixed_list_positions_are_not_ignored(self):
        with tempfile.TemporaryDirectory() as temp:
            for case in ("oneof", "mixed"):
                source = dataset()
                if case == "oneof":
                    source.reactions[0].inputs["00123"].components[0].amount.ClearField("mass")
                    source.reactions[0].inputs["00123"].components[0].amount.mass.SetInParent()
                else:
                    source.reactions[0].conditions.temperature.measurements.add()
                    source.reactions[0].conditions.temperature.measurements.add().details = "measured"
                path = Path(temp) / "source.csv"
                spreadsheet._write_csv(source, path)
                with self.subTest(case=case), self.assertRaises(ValueError):
                    spreadsheet.verify_round_trip(source, path)

    def test_changed_scientific_values_and_identity_fail_verification(self):
        source = dataset()
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "test.csv"
            spreadsheet._write_csv(source, path)
            text = path.read_text()
            for corrupted in (text.replace(DATASET_ID, "cmcc_dataset-" + "c" * 32),
                              text.replace("kneading", "changed value")):
                path.write_text(corrupted)
                with self.assertRaises(spreadsheet.SpreadsheetRoundTripError):
                    spreadsheet.verify_round_trip(source, path)

    def test_failed_csv_verification_preserves_existing_file(self):
        source = dataset(); source.reactions[0].notes.procedure_details = " boundary text "
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "existing.csv"; path.write_bytes(b"human data")
            with self.assertRaises(FileExistsError):
                spreadsheet.export_dataset(source, path)
            with self.assertRaisesRegex(ValueError, "round-trip changed"):
                spreadsheet.export_dataset(source, path, overwrite=True)
            self.assertEqual(path.read_bytes(), b"human data")


if __name__ == "__main__":
    unittest.main()
